from __future__ import annotations

import argparse
import math
import re
import sys
from pathlib import Path
from typing import Any

import psycopg
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.database import connect, initialize_database

PATTAYA_SHEET = "พัทยา ชลบุรี"
YEAR_PATTERN = re.compile(r"(?:19|20)\d{2}")
QUARTER_PATTERN = re.compile(r"(?:Q|ไตรมาส)\s*([1-4])", re.IGNORECASE)


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def numeric_value(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def find_year(value: Any) -> int | None:
    match = YEAR_PATTERN.search(clean_text(value) or "")
    return int(match.group()) if match else None


def find_quarter(value: Any) -> int | None:
    match = QUARTER_PATTERN.search(clean_text(value) or "")
    return int(match.group(1)) if match else None


def collect_sheet_cells(worksheet: Any, source_file: str) -> list[tuple[str, str, int, int, str | None]]:
    cells = []
    for row in worksheet.iter_rows():
        for cell in row:
            value = clean_text(cell.value)
            if value is not None:
                cells.append((source_file, worksheet.title, cell.row, cell.column, value))
    return cells


def extract_metrics(worksheet: Any, source_file: str) -> list[tuple[Any, ...]]:
    metrics: list[tuple[Any, ...]] = []
    active_year: int | None = None
    active_quarter: int | None = None

    for row in worksheet.iter_rows():
        values = [cell.value for cell in row]
        row_text = " | ".join(clean_text(value) or "" for value in values)
        row_years = [find_year(value) for value in values if find_year(value)]
        row_quarters = [find_quarter(value) for value in values if find_quarter(value)]
        if row_years:
            active_year = row_years[-1]
        if row_quarters:
            active_quarter = row_quarters[-1]

        labels = [clean_text(value) for value in values if clean_text(value) and numeric_value(value) is None]
        label = labels[0] if labels else None
        if not label or len(label) < 2:
            continue

        for cell in row:
            value = numeric_value(cell.value)
            if value is None:
                continue
            cell_label = label
            if cell.column > 1:
                previous = clean_text(values[cell.column - 2])
                if previous and numeric_value(previous) is None and not find_year(previous):
                    cell_label = previous
            metrics.append((
                source_file,
                worksheet.title,
                cell_label,
                value,
                "%" if "%" in row_text or "%" in cell_label else None,
                active_year,
                active_quarter,
                cell.row,
            ))
    return metrics


def import_workbook(workbook_path: Path, connection: psycopg.Connection) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    source_file = workbook_path.name
    imported_cells = 0
    imported_metrics = 0

    connection.execute("DELETE FROM raw_excel_cells WHERE source_file = %s", (source_file,))
    connection.execute("DELETE FROM pattaya_report_metrics WHERE source_file = %s", (source_file,))

    with connection.cursor() as cursor:
        for worksheet in workbook.worksheets:
            cells = collect_sheet_cells(worksheet, source_file)
            cursor.executemany(
                """
                INSERT INTO raw_excel_cells
                    (source_file, sheet_name, row_number, column_number, cell_value)
                VALUES (%s, %s, %s, %s, %s)
                """,
                cells,
            )
            imported_cells += len(cells)

            if worksheet.title == PATTAYA_SHEET:
                metrics = extract_metrics(worksheet, source_file)
                cursor.executemany(
                    """
                    INSERT INTO pattaya_report_metrics
                        (source_file, sheet_name, metric_label, metric_value, unit,
                         year, quarter, source_row)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    metrics,
                )
                imported_metrics += len(metrics)

    workbook.close()
    return {"sheets": len(workbook.sheetnames), "cells": imported_cells, "metrics": imported_metrics}


def main() -> None:
    parser = argparse.ArgumentParser(description="Import tourism Excel reports into staging tables")
    parser.add_argument("workbook", type=Path, help="Path to .xlsx workbook")
    args = parser.parse_args()
    if not args.workbook.exists():
        raise SystemExit(f"ไม่พบไฟล์: {args.workbook}")

    initialize_database()
    with connect() as connection:
        result = import_workbook(args.workbook, connection)
    print(f"นำเข้าสำเร็จ: {args.workbook.name}")
    print(f"ชีต: {result['sheets']} | cells: {result['cells']} | metrics พัทยา: {result['metrics']}")


if __name__ == "__main__":
    main()
